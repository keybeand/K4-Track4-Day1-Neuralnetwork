"""results_table.py — Quản lý lưu kết quả JSON và tự động điền vào bảng Excel.
"""
from __future__ import annotations

import json
import os
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] ra <results_dir>/<exp_id>.json."""
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    filepath = os.path.join(results_dir, f"{exp_id}.json")

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"]
    }

    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return filepath


def load_results(results_dir: str = "results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict."""
    if not os.path.exists(results_dir):
        return []

    results = []
    for file in sorted(os.listdir(results_dir)):
        if file.endswith(".json"):
            filepath = os.path.join(results_dir, file)
            with open(filepath, "r", encoding="utf-8") as f:
                results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng."""
    cfg = result["cfg"]
    summary = result["summary"]

    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg["group"],
        "description": cfg.get("description", ""),
        "loss": cfg["loss"],
        "optimizer": cfg["optimizer"],
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": str(cfg["hidden"]),
        "dropout": cfg["dropout"],
        "clip_norm": cfg["clip_norm"] if cfg["clip_norm"] is not None else "None",
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "step0_loss": summary["step0_loss"],
        "best_val_loss": summary["best_val_loss"],
        "best_epoch": summary["best_epoch"],
        "final_train_loss": summary["final_train_loss"],
        "final_val_loss": summary["final_val_loss"],
        "val_acc": summary["val_acc"],
        "val_macro_f1": summary["val_macro_f1"],
        "time_per_epoch_s": summary["time_per_epoch_s"],
        "peak_mem_MB": summary["peak_mem_MB"],
        "diverged": summary["diverged"],
        "eval_acc": eval_scores["accuracy"] if eval_scores else None,
        "eval_macro_f1": eval_scores["macro_f1"] if eval_scores else None,
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    # Đọc danh sách tên cột ở dòng 1
    headers = [cell.value for cell in ws[1]]

    # Xóa các dòng mẫu cũ từ dòng 2 nếu có
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)

    for row_idx, data_row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            if header in data_row and data_row[header] is not None:
                ws.cell(row=row_idx, column=col_idx, value=data_row[header])

    wb.save(out_path)
    print(f"Updated experiment table saved to {out_path} ({len(rows)} rows)")
